"""DOCX, XLSX, CSV parsers."""
import csv
import io
import logging
from typing import Tuple

logger = logging.getLogger(__name__)


def parse_docx(file_path: str) -> Tuple[str, int]:
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(c.text for c in row.cells if c.text.strip()))
        return "\n".join(paragraphs), len(paragraphs)
    except Exception as exc:
        logger.error("DOCX parse error %s: %s", file_path, exc)
        return "", 0


def parse_xlsx(file_path: str) -> Tuple[str, int]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            rows.append(f"=== Sheet: {sheet.title} ===")
            for row in sheet.iter_rows(values_only=True):
                clean = [str(c) if c is not None else "" for c in row]
                if any(c for c in clean):
                    rows.append(" | ".join(clean))
        return "\n".join(rows), len(rows)
    except Exception as exc:
        logger.error("XLSX parse error %s: %s", file_path, exc)
        return "", 0


def parse_csv(file_path: str) -> Tuple[str, int]:
    try:
        with open(file_path, newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.reader(f)
            rows = [" | ".join(row) for row in reader if any(c for c in row)]
        return "\n".join(rows), len(rows)
    except Exception as exc:
        logger.error("CSV parse error %s: %s", file_path, exc)
        return "", 0
