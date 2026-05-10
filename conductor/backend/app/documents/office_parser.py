"""Office document parsers: DOCX, XLSX, CSV."""
import csv
import io
from typing import Tuple
import openpyxl
from docx import Document as DocxDocument


def parse_docx(file_path: str) -> Tuple[str, int]:
    """Returns (full_text, paragraph_count)."""
    doc = DocxDocument(file_path)
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Extract tables
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            paragraphs.append("TABLE:\n" + "\n".join(rows))

    return "\n\n".join(paragraphs), len(paragraphs)


def parse_xlsx(file_path: str) -> str:
    """Convert all sheets to text."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append(" | ".join(str(c) if c is not None else "" for c in row))
        if rows:
            parts.append(f"Sheet: {sheet_name}\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def parse_csv(file_path: str) -> str:
    """Convert CSV to pipe-delimited text."""
    rows = []
    with open(file_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(" | ".join(row))
    return "\n".join(rows)
